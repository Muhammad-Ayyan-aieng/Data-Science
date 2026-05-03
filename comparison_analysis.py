# comparison_analysis.py
"""
Comparative Analysis: GA vs CT - Answering Key Questions
Identifies factors that explain state differences
Exports to: 1 Excel file, 1 Word file, multiple CSV files
"""

import pandas as pd
import numpy as np
import os
from scipy import stats
from scipy.stats import chi2_contingency, ttest_ind
from datetime import datetime
import warnings
warnings.filterwarnings('ignore')

# Optional imports for Excel and Word
try:
    from openpyxl import Workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("⚠️ openpyxl not installed. Excel export limited.")

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False
    print("⚠️ python-docx not installed. Word export limited.")

def run_comparison_analysis():
    """Analyze WHY Georgia has different severity than Connecticut"""
    
    print("=" * 80)
    print("🔍 COMPARATIVE ANALYSIS: GA vs CT - WHY THE DIFFERENCE?")
    print("=" * 80)
    
    # Create output folders
    os.makedirs("outputs/comparison_analysis/tables", exist_ok=True)
    os.makedirs("outputs/comparison_analysis/reports", exist_ok=True)
    
    # Load data
    input_file = os.path.join("datasets", "cleaned_accidents.csv")
    df = pd.read_csv(input_file)
    print(f"\n✅ Loaded {len(df):,} accidents")
    
    ga_df = df[df['State'] == 'GA']
    ct_df = df[df['State'] == 'CT']
    
    print(f"\n📊 Georgia: {len(ga_df):,} accidents (Severity: {ga_df['Severity'].mean():.3f})")
    print(f"📊 Connecticut: {len(ct_df):,} accidents (Severity: {ct_df['Severity'].mean():.3f})")
    print(f"📊 Severity Difference: {ga_df['Severity'].mean() - ct_df['Severity'].mean():+.3f}")
    
    # Dictionary to store all tables for Excel export
    all_tables = {}
    
    # =========================================================
    # QUESTION 1: Is the severity difference statistically significant?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 1: Is the severity difference statistically significant?")
    print("=" * 80)
    
    t_stat, p_value = ttest_ind(ga_df['Severity'], ct_df['Severity'])
    
    print(f"\nT-test results:")
    print(f"   T-statistic: {t_stat:.4f}")
    print(f"   P-value: {p_value:.6f}")
    
    if p_value < 0.05:
        print(f"   ✅ Difference is STATISTICALLY SIGNIFICANT (p < 0.05)")
        print(f"   → Georgia's higher severity is REAL, not random chance")
    else:
        print(f"   ❌ Difference is NOT statistically significant")
    
    significance_df = pd.DataFrame([{
        'Metric': 'Severity Difference',
        'Georgia_Mean': ga_df['Severity'].mean(),
        'Connecticut_Mean': ct_df['Severity'].mean(),
        'Difference': ga_df['Severity'].mean() - ct_df['Severity'].mean(),
        'T_Statistic': t_stat,
        'P_Value': p_value,
        'Significant': p_value < 0.05
    }])
    all_tables['Statistical_Significance'] = significance_df
    
    # =========================================================
    # QUESTION 2: How do weather conditions differ between states?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 2: How do weather conditions differ between states?")
    print("=" * 80)
    
    weather_features = ['Temperature(F)', 'Humidity(%)', 'Precipitation(in)', 'Visibility(mi)']
    weather_comparison = []
    
    for feat in weather_features:
        if feat in df.columns:
            ga_mean = ga_df[feat].mean()
            ct_mean = ct_df[feat].mean()
            t_stat_w, p_val_w = ttest_ind(ga_df[feat].dropna(), ct_df[feat].dropna())
            
            weather_comparison.append({
                'Weather_Factor': feat,
                'Georgia': round(ga_mean, 2),
                'Connecticut': round(ct_mean, 2),
                'Difference': round(ga_mean - ct_mean, 2),
                'P_Value': round(p_val_w, 4),
                'Significant': p_val_w < 0.05
            })
    
    weather_df = pd.DataFrame(weather_comparison)
    print("\n", weather_df.to_string(index=False))
    all_tables['Weather_Comparison'] = weather_df
    
    # =========================================================
    # QUESTION 3: What time of day do accidents happen in each state?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 3: What time of day do accidents happen in each state?")
    print("=" * 80)
    
    # Initialize variables with defaults
    ga_peak = "N/A"
    ct_peak = "N/A"
    has_timeofday = False
    
    if 'TimeOfDay' in df.columns:
        has_timeofday = True
        time_order = ['Late Night', 'Morning', 'Afternoon', 'Evening', 'Night']
        ga_time = ga_df['TimeOfDay'].value_counts(normalize=True) * 100
        ct_time = ct_df['TimeOfDay'].value_counts(normalize=True) * 100
        
        time_comparison = []
        for t in time_order:
            ga_pct = ga_time.get(t, 0)
            ct_pct = ct_time.get(t, 0)
            time_comparison.append({
                'Time_Period': t,
                'Georgia_%': round(ga_pct, 1),
                'Connecticut_%': round(ct_pct, 1),
                'Difference': round(ga_pct - ct_pct, 1)
            })
        
        time_df = pd.DataFrame(time_comparison)
        print("\n", time_df.to_string(index=False))
        all_tables['Time_of_Day_Comparison'] = time_df
        
        # Get peak accident times
        ga_peak = time_df.loc[time_df['Georgia_%'].idxmax(), 'Time_Period']
        ct_peak = time_df.loc[time_df['Connecticut_%'].idxmax(), 'Time_Period']
        
        # Chi-square test for time distribution difference
        ga_time_counts = ga_df['TimeOfDay'].value_counts()
        ct_time_counts = ct_df['TimeOfDay'].value_counts()
        contingency = pd.DataFrame([ga_time_counts, ct_time_counts]).fillna(0)
        chi2, p_val_chi, dof, expected = chi2_contingency(contingency)
        
        print(f"\nChi-square test for time distribution:")
        print(f"   Chi-square: {chi2:.2f}")
        print(f"   P-value: {p_val_chi:.4f}")
        print(f"   Time distributions are {'DIFFERENT' if p_val_chi < 0.05 else 'SIMILAR'} between states")
    else:
        print("\n⚠️ TimeOfDay column not found. Skipping time analysis.")
        time_df = pd.DataFrame()  # Empty DataFrame
    
    # =========================================================
    # QUESTION 4: Which days are most dangerous in each state?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 4: Which days are most dangerous in each state?")
    print("=" * 80)
    
    if 'DayOfWeek' in df.columns:
        day_names = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
        ga_day = ga_df.groupby('DayOfWeek')['Severity'].mean()
        ct_day = ct_df.groupby('DayOfWeek')['Severity'].mean()
        
        day_comparison = []
        for i, day in enumerate(day_names):
            ga_sev = ga_day.get(i, 0)
            ct_sev = ct_day.get(i, 0)
            day_comparison.append({
                'Day': day,
                'Georgia_Severity': round(ga_sev, 3),
                'Connecticut_Severity': round(ct_sev, 3),
                'Difference': round(ga_sev - ct_sev, 3)
            })
        
        day_df = pd.DataFrame(day_comparison)
        print("\n", day_df.to_string(index=False))
        all_tables['Day_of_Week_Comparison'] = day_df
    else:
        print("\n⚠️ DayOfWeek column not found. Skipping day analysis.")
        day_df = pd.DataFrame()
    
    # =========================================================
    # QUESTION 5: How do road features differ between states?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 5: How do road features differ between states?")
    print("=" * 80)
    
    road_features = ['Junction', 'Traffic_Signal', 'Crossing', 'Railway', 
                     'Stop', 'Traffic_Calming', 'Roundabout', 'Amenity']
    
    road_comparison = []
    for feat in road_features:
        if feat in df.columns:
            ga_pct = (ga_df[feat].sum() / len(ga_df)) * 100
            ct_pct = (ct_df[feat].sum() / len(ct_df)) * 100
            
            # Chi-square test for difference
            contingency_road = pd.crosstab(df['State'], df[feat])
            chi2_r, p_val_r, dof_r, expected_r = chi2_contingency(contingency_road)
            
            road_comparison.append({
                'Road_Feature': feat,
                'Georgia_%': round(ga_pct, 1),
                'Connecticut_%': round(ct_pct, 1),
                'Difference': round(ga_pct - ct_pct, 1),
                'P_Value': round(p_val_r, 4),
                'Significant': p_val_r < 0.05
            })
    
    road_feat_df = pd.DataFrame(road_comparison)
    print("\n", road_feat_df.to_string(index=False))
    all_tables['Road_Features_Comparison'] = road_feat_df
    
    # =========================================================
    # QUESTION 6: Which cities are most dangerous in each state?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 6: Most dangerous cities in each state")
    print("=" * 80)
    
    # Georgia's most dangerous cities
    ga_cities = ga_df.groupby('City')['Severity'].agg(['mean', 'count']).reset_index()
    ga_cities = ga_cities[ga_cities['count'] >= 50].sort_values('mean', ascending=False).head(5)
    ga_cities.columns = ['City', 'Avg_Severity', 'Accident_Count']
    
    # Connecticut's most dangerous cities
    ct_cities = ct_df.groupby('City')['Severity'].agg(['mean', 'count']).reset_index()
    ct_cities = ct_cities[ct_cities['count'] >= 50].sort_values('mean', ascending=False).head(5)
    ct_cities.columns = ['City', 'Avg_Severity', 'Accident_Count']
    
    print("\n🔴 GEORGIA'S MOST DANGEROUS CITIES:")
    print(ga_cities.to_string(index=False))
    print("\n🔴 CONNECTICUT'S MOST DANGEROUS CITIES:")
    print(ct_cities.to_string(index=False))
    
    all_tables['GA_Most_Dangerous_Cities'] = ga_cities
    all_tables['CT_Most_Dangerous_Cities'] = ct_cities
    
    # =========================================================
    # QUESTION 7: Which factors contribute most to severity difference?
    # =========================================================
    print("\n" + "=" * 80)
    print("QUESTION 7: What factors explain the severity difference?")
    print("=" * 80)
    
    # Compare severity when same conditions exist
    contributing_factors = []
    
    # Compare at same time of day
    if 'TimeOfDay' in df.columns:
        for time in ['Morning', 'Afternoon', 'Evening', 'Night']:
            ga_time_sev = ga_df[ga_df['TimeOfDay'] == time]['Severity'].mean()
            ct_time_sev = ct_df[ct_df['TimeOfDay'] == time]['Severity'].mean()
            if not pd.isna(ga_time_sev) and not pd.isna(ct_time_sev):
                contributing_factors.append({
                    'Condition': f'{time} accidents',
                    'Georgia_Severity': round(ga_time_sev, 3),
                    'Connecticut_Severity': round(ct_time_sev, 3),
                    'Difference': round(ga_time_sev - ct_time_sev, 3)
                })
    
    # Compare at same road features
    for feat in ['Junction', 'Traffic_Signal', 'Crossing']:
        if feat in df.columns:
            ga_feat_sev = ga_df[ga_df[feat] == True]['Severity'].mean()
            ct_feat_sev = ct_df[ct_df[feat] == True]['Severity'].mean()
            if not pd.isna(ga_feat_sev) and not pd.isna(ct_feat_sev):
                contributing_factors.append({
                    'Condition': f'At {feat}',
                    'Georgia_Severity': round(ga_feat_sev, 3),
                    'Connecticut_Severity': round(ct_feat_sev, 3),
                    'Difference': round(ga_feat_sev - ct_feat_sev, 3)
                })
    
    factors_df = pd.DataFrame(contributing_factors)
    print("\n", factors_df.to_string(index=False))
    all_tables['Contributing_Factors'] = factors_df
    
    # =========================================================
    # SAVE CSV FILES
    # =========================================================
    print("\n" + "-" * 50)
    print("SAVING CSV FILES")
    print("-" * 50)
    
    significance_df.to_csv("outputs/comparison_analysis/tables/01_significance.csv", index=False)
    weather_df.to_csv("outputs/comparison_analysis/tables/02_weather_comparison.csv", index=False)
    
    if has_timeofday and len(time_df) > 0:
        time_df.to_csv("outputs/comparison_analysis/tables/03_time_comparison.csv", index=False)
    
    if len(day_df) > 0:
        day_df.to_csv("outputs/comparison_analysis/tables/04_day_comparison.csv", index=False)
    
    road_feat_df.to_csv("outputs/comparison_analysis/tables/05_road_features_comparison.csv", index=False)
    ga_cities.to_csv("outputs/comparison_analysis/tables/06_ga_dangerous_cities.csv", index=False)
    ct_cities.to_csv("outputs/comparison_analysis/tables/07_ct_dangerous_cities.csv", index=False)
    factors_df.to_csv("outputs/comparison_analysis/tables/08_contributing_factors.csv", index=False)
    
    print("   ✅ Saved CSV files to outputs/comparison_analysis/tables/")
    
    # =========================================================
    # CREATE EXCEL FILE
    # =========================================================
    print("\n" + "-" * 50)
    print("CREATING EXCEL FILE")
    print("-" * 50)
    
    excel_path = "outputs/comparison_analysis/reports/comparison_analysis.xlsx"
    
    if HAS_OPENPYXL:
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            for sheet_name, table in all_tables.items():
                safe_name = sheet_name.replace(' ', '_')[:31]
                table.to_excel(writer, sheet_name=safe_name, index=False)
            
            # Add info sheet
            info_df = pd.DataFrame({
                'Information': ['Analysis Date', 'Total Accidents', 'Georgia Accidents', 'Connecticut Accidents'],
                'Value': [datetime.now().strftime('%Y-%m-%d %H:%M:%S'), len(df), len(ga_df), len(ct_df)]
            })
            info_df.to_excel(writer, sheet_name='Info', index=False)
        
        print(f"   ✅ Excel file saved: {excel_path}")
    else:
        print("   ⚠️ openpyxl not installed. Skipping Excel export.")
    
    # =========================================================
    # CREATE WORD REPORT
    # =========================================================
    print("\n" + "-" * 50)
    print("CREATING WORD REPORT")
    print("-" * 50)
    
    word_path = "outputs/comparison_analysis/reports/comparison_analysis.docx"
    
    if HAS_PYTHON_DOCX:
        doc = Document()
        
        # Title
        title = doc.add_heading('GA vs CT Comparative Analysis Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Info
        doc.add_paragraph(f"Analysis Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()
        
        # Executive Summary
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph(f"""
Georgia has {len(ga_df):,} accidents ({len(ga_df)/len(df)*100:.1f}% of total) with average severity {ga_df['Severity'].mean():.2f}.
Connecticut has {len(ct_df):,} accidents ({len(ct_df)/len(df)*100:.1f}% of total) with average severity {ct_df['Severity'].mean():.2f}.
The severity difference of {ga_df['Severity'].mean() - ct_df['Severity'].mean():+.2f} is {'statistically significant' if p_value < 0.05 else 'not statistically significant'}.
""")
        
        # Key Findings
        doc.add_heading('Key Findings', level=1)
        
        # Weather difference
        if len(weather_df) > 0:
            temp_row = weather_df[weather_df['Weather_Factor'] == 'Temperature(F)']
            if len(temp_row) > 0:
                temp_diff = temp_row['Difference'].values[0]
                doc.add_paragraph(f"• Weather: Georgia is {abs(temp_diff):.1f}°F {'warmer' if temp_diff > 0 else 'colder'} than Connecticut on average")
        
        # Time difference
        if has_timeofday and ga_peak != "N/A" and ct_peak != "N/A":
            doc.add_paragraph(f"• Peak accident time: {ga_peak} in Georgia vs {ct_peak} in Connecticut")
        
        # Most dangerous cities
        if len(ga_cities) > 0:
            doc.add_heading('Most Dangerous Cities', level=2)
            doc.add_paragraph(f"Georgia's most dangerous city: {ga_cities.iloc[0]['City']} (Severity: {ga_cities.iloc[0]['Avg_Severity']:.2f})")
            doc.add_paragraph(f"Connecticut's most dangerous city: {ct_cities.iloc[0]['City']} (Severity: {ct_cities.iloc[0]['Avg_Severity']:.2f})")
        
        # Weather Comparison Table
        if len(weather_df) > 0:
            doc.add_heading('Weather Comparison', level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            hdr[0].text = 'Weather Factor'
            hdr[1].text = 'Georgia'
            hdr[2].text = 'Connecticut'
            hdr[3].text = 'Difference'
            
            for _, row in weather_df.iterrows():
                cells = table.add_row().cells
                cells[0].text = row['Weather_Factor']
                cells[1].text = str(row['Georgia'])
                cells[2].text = str(row['Connecticut'])
                cells[3].text = str(row['Difference'])
            
            doc.add_paragraph()
        
        # Time Comparison Table
        if has_timeofday and len(time_df) > 0:
            doc.add_heading('Time of Day Comparison', level=1)
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            hdr[0].text = 'Time Period'
            hdr[1].text = 'Georgia %'
            hdr[2].text = 'Connecticut %'
            
            for _, row in time_df.iterrows():
                cells = table.add_row().cells
                cells[0].text = row['Time_Period']
                cells[1].text = str(row['Georgia_%'])
                cells[2].text = str(row['Connecticut_%'])
        
        # Save Word document
        doc.save(word_path)
        print(f"   ✅ Word report saved: {word_path}")
    else:
        print("   ⚠️ python-docx not installed. Skipping Word export.")
    
    # =========================================================
    # FINAL SUMMARY
    # =========================================================
    print("\n" + "=" * 80)
    print("📊 COMPARATIVE ANALYSIS - SUMMARY OF FINDINGS")
    print("=" * 80)
    
    # Safely get values for summary
    temp_diff_val = 0
    if len(weather_df) > 0:
        temp_row = weather_df[weather_df['Weather_Factor'] == 'Temperature(F)']
        if len(temp_row) > 0:
            temp_diff_val = temp_row['Difference'].values[0]
    
    print(f"""
✅ ANALYSIS COMPLETE!

KEY QUESTIONS ANSWERED:
---------------------
1. Is severity difference significant? 
   → {'YES ✅' if p_value < 0.05 else 'NO ❌'} (p={p_value:.4f})

2. Weather differences:
   → Georgia is {temp_diff_val:+.1f}°F {'warmer' if temp_diff_val > 0 else 'colder'}

3. Peak accident time:
   → Georgia: {ga_peak}, Connecticut: {ct_peak}

4. Most dangerous city:
   → Georgia: {ga_cities.iloc[0]['City'] if len(ga_cities) > 0 else 'N/A'} (Severity: {ga_cities.iloc[0]['Avg_Severity'] if len(ga_cities) > 0 else 'N/A'})
   → Connecticut: {ct_cities.iloc[0]['City'] if len(ct_cities) > 0 else 'N/A'} (Severity: {ct_cities.iloc[0]['Avg_Severity'] if len(ct_cities) > 0 else 'N/A'})

OUTPUT FILES:
------------
📁 CSV Files: outputs/comparison_analysis/tables/ (8 files)
📊 Excel File: outputs/comparison_analysis/reports/comparison_analysis.xlsx
📄 Word Report: outputs/comparison_analysis/reports/comparison_analysis.docx
""")
    
    print("=" * 80)
    print("✨ COMPARATIVE ANALYSIS COMPLETE! ✨")
    print("=" * 80)
    
    return all_tables

if __name__ == "__main__" or __name__ == "comparison_analysis":
    run_comparison_analysis()